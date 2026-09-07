from __future__ import annotations

import hashlib
import re
from pathlib import Path
from typing import Any, Iterable

from .canonical_compatibility import V1_SOURCE_ORDER
from .canonical_materialization_verify import expected_permit_build_id
from .storage import is_within, resolve_data_root


REFERENCE_SHEET_NAME = "1. 개방자치단체코드"
REFERENCE_HEADERS = (
    "번호",
    "자치단체명",
    "자치단체 코드",
    "변경여부",
    "변경일",
    "변경사유",
)
REFERENCE_ATTACHMENT_FILENAME = "개방자치단체코드_영업상태코드_행정체제개편반영_20260702.xlsx"
REFERENCE_NOTICE_URL = (
    "https://www.data.go.kr/bbs/ntc/selectNotice.do?originId=NOTICE_0000000004824"
)
REFERENCE_NOTICE_DATE = "2026-07-03"
REFERENCE_CHANGE_EFFECTIVE_DATE = "2026-07-01"
REFERENCE_CHANGE_EFFECTIVE_DATE_BASIC = "20260701"
MANUAL_URL = "https://www.localdata.go.kr/images/egovframework/portal/manual_260106.pdf"
MANUAL_REFERENCE_COUNT = 245
V1_DETAIL_REFERENCE_FILENAME = "개방자치단체코드_영업상태코드.xlsx"
V1_DETAIL_PAGES = (
    "https://www.data.go.kr/data/15154916/openapi.do",
    "https://www.data.go.kr/data/15154921/openapi.do",
    "https://www.data.go.kr/data/15155252/openapi.do",
)
PYARROW_VERSION = "21.0.0"

_NUMERIC_CODE = re.compile(r"^[0-9]{7}$")
_AGGREGATE_CODE = re.compile(r"^[0-9]{7}_ALL$")


class AuthorityDomainError(RuntimeError):
    """Raised when official authority-domain evidence cannot be parsed safely."""


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _normalize_number(value: object) -> int | None:
    if value is None:
        return None
    if isinstance(value, bool):
        raise AuthorityDomainError("authority reference row number cannot be boolean")
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    raise AuthorityDomainError("authority reference row number must be an integer or blank")


def _normalize_text(value: object, *, field: str, allow_blank: bool = False) -> str | None:
    if value is None:
        if allow_blank:
            return None
        raise AuthorityDomainError(f"authority reference {field} cannot be blank")
    text = str(value).strip()
    if not text:
        if allow_blank:
            return None
        raise AuthorityDomainError(f"authority reference {field} cannot be blank")
    return text


def _normalize_code(value: object) -> str:
    if isinstance(value, bool):
        raise AuthorityDomainError("authority code cannot be boolean")
    if isinstance(value, int):
        code = str(value)
    elif isinstance(value, float) and value.is_integer():
        code = str(int(value))
    else:
        code = _normalize_text(value, field="code") or ""
    if not (_NUMERIC_CODE.fullmatch(code) or _AGGREGATE_CODE.fullmatch(code)):
        raise AuthorityDomainError("authority reference contains an unsupported code format")
    return code


def _normalize_change_date(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        raise AuthorityDomainError("authority reference change date cannot be boolean")
    if isinstance(value, int):
        raw = str(value)
    elif isinstance(value, float) and value.is_integer():
        raw = str(int(value))
    else:
        raw = str(value).strip().replace("-", "")
    if not re.fullmatch(r"[0-9]{8}", raw):
        raise AuthorityDomainError("authority reference change date must be YYYYMMDD or blank")
    return raw


def _authority_list_hash(authorities: Iterable[dict[str, str]]) -> str:
    lines = [f"{item['code']}\t{item['name']}\n" for item in authorities]
    return hashlib.sha256("".join(lines).encode("utf-8")).hexdigest()


def parse_authority_reference_workbook(path: str | Path) -> dict[str, Any]:
    """Parse the official current authority-code workbook without source business rows."""

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise AuthorityDomainError(
            'openpyxl is required for authority-reference parsing; install with '
            'python -m pip install -e ".[reference]"'
        ) from exc

    artifact = Path(path).expanduser().resolve()
    if not artifact.is_file():
        raise AuthorityDomainError("authority reference workbook does not exist")

    workbook = load_workbook(artifact, read_only=True, data_only=True)
    try:
        if REFERENCE_SHEET_NAME not in workbook.sheetnames:
            raise AuthorityDomainError("authority reference sheet name changed")
        sheet = workbook[REFERENCE_SHEET_NAME]
        headers = tuple(sheet.cell(1, column).value for column in range(1, 7))
        if headers != REFERENCE_HEADERS:
            raise AuthorityDomainError("authority reference headers changed")

        active_rows: list[dict[str, Any]] = []
        deleted_rows: list[dict[str, Any]] = []
        for values in sheet.iter_rows(min_row=2, max_col=6, values_only=True):
            if all(value is None for value in values):
                continue
            number = _normalize_number(values[0])
            name = _normalize_text(values[1], field="name") or ""
            code = _normalize_code(values[2])
            change = _normalize_text(values[3], field="change", allow_blank=True)
            change_date = _normalize_change_date(values[4])

            row = {
                "number": number,
                "name": name,
                "code": code,
                "change": change,
                "change_date": change_date,
            }
            if number is not None:
                if change not in {None, "신규"}:
                    raise AuthorityDomainError("numbered authority row has an invalid change marker")
                if change is None and change_date is not None:
                    raise AuthorityDomainError("unchanged authority row unexpectedly has a change date")
                if change == "신규" and _NUMERIC_CODE.fullmatch(code):
                    if change_date != REFERENCE_CHANGE_EFFECTIVE_DATE_BASIC:
                        raise AuthorityDomainError("new numeric authority row effective date changed")
                if change == "신규" and _AGGREGATE_CODE.fullmatch(code):
                    if change_date is not None:
                        raise AuthorityDomainError("new aggregate authority token unexpectedly has a change date")
                active_rows.append(row)
            else:
                if change != "삭제":
                    raise AuthorityDomainError("unnumbered authority row is not an explicit deletion")
                if change_date != REFERENCE_CHANGE_EFFECTIVE_DATE_BASIC:
                    raise AuthorityDomainError("deleted authority row effective date changed")
                deleted_rows.append(row)
    finally:
        workbook.close()

    numbers = [int(item["number"]) for item in active_rows]
    if numbers != list(range(1, len(active_rows) + 1)):
        raise AuthorityDomainError("active authority row numbers are not contiguous")

    active_codes = [str(item["code"]) for item in active_rows]
    if len(active_codes) != len(set(active_codes)):
        raise AuthorityDomainError("active authority codes are not unique")
    active_names = [str(item["name"]) for item in active_rows]
    if len(active_names) != len(set(active_names)):
        raise AuthorityDomainError("active authority names are not unique")

    active_numeric = sorted(
        (
            {"code": str(item["code"]), "name": str(item["name"])}
            for item in active_rows
            if _NUMERIC_CODE.fullmatch(str(item["code"]))
        ),
        key=lambda item: item["code"],
    )
    deleted_numeric = sorted(
        (
            {"code": str(item["code"]), "name": str(item["name"])}
            for item in deleted_rows
            if _NUMERIC_CODE.fullmatch(str(item["code"]))
        ),
        key=lambda item: item["code"],
    )
    new_numeric = sorted(
        (
            {"code": str(item["code"]), "name": str(item["name"])}
            for item in active_rows
            if item["change"] == "신규" and _NUMERIC_CODE.fullmatch(str(item["code"]))
        ),
        key=lambda item: item["code"],
    )
    active_aggregate = sorted(
        (
            {"code": str(item["code"]), "name": str(item["name"])}
            for item in active_rows
            if _AGGREGATE_CODE.fullmatch(str(item["code"]))
        ),
        key=lambda item: item["code"],
    )
    new_numeric_codes = {
        str(item["code"])
        for item in active_rows
        if item["change"] == "신규" and _NUMERIC_CODE.fullmatch(str(item["code"]))
    }
    new_aggregate_codes = {
        str(item["code"])
        for item in active_rows
        if item["change"] == "신규" and _AGGREGATE_CODE.fullmatch(str(item["code"]))
    }
    deleted_numeric_codes = {
        str(item["code"])
        for item in deleted_rows
        if _NUMERIC_CODE.fullmatch(str(item["code"]))
    }
    deleted_aggregate_codes = {
        str(item["code"])
        for item in deleted_rows
        if _AGGREGATE_CODE.fullmatch(str(item["code"]))
    }

    return {
        "artifact_sha256": _sha256_file(artifact),
        "sheet_name": REFERENCE_SHEET_NAME,
        "headers": list(REFERENCE_HEADERS),
        "active_row_count": len(active_rows),
        "active_numeric_code_count": len(active_numeric),
        "active_aggregate_token_count": len(active_aggregate),
        "deleted_row_count": len(deleted_rows),
        "deleted_numeric_code_count": len(deleted_numeric_codes),
        "deleted_aggregate_token_count": len(deleted_aggregate_codes),
        "new_row_count": sum(item["change"] == "신규" for item in active_rows),
        "new_numeric_code_count": len(new_numeric_codes),
        "new_aggregate_token_count": len(new_aggregate_codes),
        "active_numeric_authorities": active_numeric,
        "active_numeric_authority_list_sha256": _authority_list_hash(active_numeric),
        "deleted_numeric_authorities": deleted_numeric,
        "deleted_numeric_authority_list_sha256": _authority_list_hash(deleted_numeric),
        "new_numeric_authorities": new_numeric,
        "new_numeric_authority_list_sha256": _authority_list_hash(new_numeric),
        "active_aggregate_tokens": active_aggregate,
        "new_numeric_codes": sorted(new_numeric_codes),
        "deleted_numeric_codes": sorted(deleted_numeric_codes),
    }


def observed_authority_sets_from_permit_build(build_dir: str | Path) -> dict[str, set[str]]:
    """Read only the canonical authority-code column and return one set per v1 source."""

    try:
        import pyarrow
        import pyarrow.compute as pc
        import pyarrow.parquet as pq
    except ImportError as exc:
        raise AuthorityDomainError(
            'pyarrow is required for current-domain comparison; install with '
            'python -m pip install -e ".[build]"'
        ) from exc
    if pyarrow.__version__ != PYARROW_VERSION:
        raise AuthorityDomainError(
            f"authority-domain comparison requires pyarrow {PYARROW_VERSION}, "
            f"found {pyarrow.__version__}"
        )

    root = Path(build_dir).expanduser().resolve()
    if not root.is_dir():
        raise AuthorityDomainError("canonical PERMIT build directory does not exist")

    result: dict[str, set[str]] = {}
    for source_key in V1_SOURCE_ORDER:
        parquet_path = root / f"{source_key}.parquet"
        if not parquet_path.is_file():
            raise AuthorityDomainError(f"{source_key}: canonical PERMIT parquet is missing")
        values = pc.unique(pq.read_table(parquet_path, columns=["authority_code"])["authority_code"])
        codes = {str(value.as_py()) for value in values if value.as_py() is not None}
        if any(_NUMERIC_CODE.fullmatch(code) is None for code in codes):
            raise AuthorityDomainError(f"{source_key}: observed authority-code format changed")
        result[source_key] = codes
    return result


def build_authority_domain_reference(
    *,
    reference_xlsx: str | Path,
    data_root: str | Path | None = None,
    build_id: str | None = None,
) -> dict[str, Any]:
    """Build safe tracked evidence from the official workbook and verified local PERMIT build."""

    root = resolve_data_root(data_root)
    reference_path = Path(reference_xlsx).expanduser().resolve()
    if not is_within(reference_path, root):
        raise AuthorityDomainError("official authority workbook must remain under KBL_DATA_ROOT")

    expected_build = expected_permit_build_id()
    selected_build = build_id or expected_build
    if selected_build != expected_build:
        raise AuthorityDomainError("requested PERMIT build does not match the approved deterministic build")
    build_dir = root / "canonical" / "permit" / "v1" / selected_build
    if not is_within(build_dir, root):
        raise AuthorityDomainError("canonical PERMIT build escaped KBL_DATA_ROOT")

    workbook = parse_authority_reference_workbook(reference_path)
    observed_by_source = observed_authority_sets_from_permit_build(build_dir)
    observed_sets = list(observed_by_source.values())
    identical = len({frozenset(values) for values in observed_sets}) == 1
    observed = set().union(*observed_sets)
    current_map = {
        str(item["code"]): str(item["name"])
        for item in workbook["active_numeric_authorities"]
    }
    deleted_map = {
        str(item["code"]): str(item["name"])
        for item in workbook["deleted_numeric_authorities"]
    }
    new_map = {
        str(item["code"]): str(item["name"])
        for item in workbook["new_numeric_authorities"]
    }
    official = set(current_map)
    deleted = set(deleted_map)
    new = set(new_map)
    if official & deleted:
        raise AuthorityDomainError("current and deleted authority-code domains overlap")
    if not new <= official:
        raise AuthorityDomainError("new authority codes must be a subset of the current official domain")
    if new & deleted:
        raise AuthorityDomainError("new and deleted authority-code domains overlap")
    if len(new) != len(deleted):
        raise AuthorityDomainError("numeric new/deleted authority counts must balance at the reform boundary")
    unchanged = official - new
    pre_reform = unchanged | deleted
    if len(pre_reform) != len(official):
        raise AuthorityDomainError("derived pre-reform numeric authority count changed")
    history_window_candidate = official | deleted
    observed_not_official = sorted(observed - official)
    if observed_not_official:
        raise AuthorityDomainError("current snapshots contain authority codes absent from official reference")
    missing = sorted(official - observed)
    new_numeric_codes = set(workbook["new_numeric_codes"])

    return {
        "checked_at": "2026-09-07",
        "decision": "DATE_EFFECTIVE_CURRENT_STATE_AUTHORITY_DOMAINS_APPROVED_PRE_244_POST_244",
        "field": "OPN_ATMY_GRP_CD",
        "official_reference": {
            "manual_url": MANUAL_URL,
            "manual_reference_filename": "개방자치단체코드.xlsx",
            "manual_claim_count": MANUAL_REFERENCE_COUNT,
            "manual_claim_is_current_enumeration_count": False,
            "v1_detail_reference_filename": V1_DETAIL_REFERENCE_FILENAME,
            "v1_detail_pages": list(V1_DETAIL_PAGES),
            "notice_url": REFERENCE_NOTICE_URL,
            "notice_date": REFERENCE_NOTICE_DATE,
            "change_effective_date": REFERENCE_CHANGE_EFFECTIVE_DATE,
            "attachment_filename": REFERENCE_ATTACHMENT_FILENAME,
            "attachment_sha256": workbook["artifact_sha256"],
            "sheet_name": workbook["sheet_name"],
            "active_row_count": workbook["active_row_count"],
            "active_numeric_code_count": workbook["active_numeric_code_count"],
            "active_aggregate_token_count": workbook["active_aggregate_token_count"],
            "deleted_row_count": workbook["deleted_row_count"],
            "deleted_numeric_code_count": workbook["deleted_numeric_code_count"],
            "deleted_aggregate_token_count": workbook["deleted_aggregate_token_count"],
            "new_row_count": workbook["new_row_count"],
            "new_numeric_code_count": workbook["new_numeric_code_count"],
            "new_aggregate_token_count": workbook["new_aggregate_token_count"],
            "current_numeric_authority_list_sha256": workbook[
                "active_numeric_authority_list_sha256"
            ],
            "deleted_numeric_authority_list_sha256": workbook[
                "deleted_numeric_authority_list_sha256"
            ],
            "new_numeric_authority_list_sha256": workbook[
                "new_numeric_authority_list_sha256"
            ],
            "manual_count_differs_from_current_numeric_reference": (
                MANUAL_REFERENCE_COUNT != workbook["active_numeric_code_count"]
            ),
        },
        "history_window_change_reference": {
            "window_start": "2026-01-01",
            "window_end": "2026-09-06",
            "change_effective_date": REFERENCE_CHANGE_EFFECTIVE_DATE,
            "exact_deleted_numeric_authority_codes_ingested": True,
            "deleted_numeric_authority_count": len(deleted),
            "exact_new_numeric_authority_codes_ingested": True,
            "new_numeric_authority_count": len(new),
            "unchanged_numeric_authority_count": len(unchanged),
            "pre_reform_numeric_authority_count": len(pre_reform),
            "pre_reform_numeric_authority_list_sha256": _authority_list_hash(
                sorted(
                    (
                        {"code": code, "name": (current_map | deleted_map)[code]}
                        for code in pre_reform
                    ),
                    key=lambda item: item["code"],
                )
            ),
            "pre_reform_current_state_enumeration_policy": "CURRENT_MINUS_NEW_PLUS_DELETED",
            "post_reform_current_state_enumeration_policy": "CURRENT_ONLY_EXCLUDE_DELETED",
            "authority_domain_switch_uses_official_change_effective_date": True,
            "api_queryability_does_not_define_date_effective_membership": True,
            "current_plus_deleted_candidate_union_count": len(history_window_candidate),
            "current_plus_deleted_candidate_union_sha256": _authority_list_hash(
                sorted(
                    (
                        {"code": code, "name": (current_map | deleted_map)[code]}
                        for code in history_window_candidate
                    ),
                    key=lambda item: item["code"],
                )
            ),
            "date_effective_history_authority_filter_semantics_verified": False,
            "date_effective_current_state_enumeration_policy_approved": True,
            "deleted_codes_confirmed_queryable_for_prechange_base_dates": True,
            "manual_245_count_reconciled_to_exact_window_domain": False,
        },
        "observed_current_snapshots": {
            "permit_build_id": selected_build,
            "distinct_authority_count": len(observed),
            "identical_across_three_v1_sources": identical,
            "format": "7-digit numeric strings",
            "observed_is_subset_of_current_official_numeric_domain": observed <= official,
            "observed_not_in_current_official_numeric_domain_count": len(observed_not_official),
            "current_official_numeric_not_observed_count": len(missing),
            "observed_new_numeric_code_count": len(observed & new_numeric_codes),
            "current_official_numeric_not_observed": [
                {"code": code, "name": current_map[code]} for code in missing
            ],
        },
        "ingestion_gate": {
            "exact_current_official_numeric_code_values_ingested": True,
            "exact_current_official_numeric_code_list_hash_recorded": True,
            "exact_current_official_numeric_code_list_validated": True,
            "current_official_numeric_domain_authoritatively_complete_for_reference_date": True,
            "current_reference_numeric_enumeration_ready": True,
            "exact_deleted_numeric_code_values_ingested": True,
            "exact_new_numeric_code_values_ingested": True,
            "history_window_date_effective_numeric_enumeration_ready": True,
            "aggregate_all_tokens_used_for_row_enumeration": False,
            "future_reference_refresh_required": True,
        },
        "cost_model_policy": {
            "observed_nonempty_authority_count": len(observed),
            "current_official_numeric_authority_count": len(official),
            "current_unobserved_official_authority_probe_count_per_source": len(missing),
            "deleted_numeric_authority_count": len(deleted),
            "new_numeric_authority_count": len(new),
            "pre_reform_numeric_authority_count": len(pre_reform),
            "history_window_candidate_numeric_union_count": len(history_window_candidate),
            "current_scale_candidate_union_absent_authority_count_per_source": (
                len(history_window_candidate) - len(observed)
            ),
            "current_plus_deleted_candidate_union_may_be_used_for_cost_planning": False,
            "candidate_union_is_proven_date_effective_query_domain": False,
            "date_effective_current_state_authority_count_per_date": len(official),
            "current_unobserved_authorities_may_be_assumed_historically_empty": False,
            "deleted_authorities_may_be_assumed_unqueryable_for_prechange_dates": False,
            "manual_245_count_may_be_used_as_exact_current_query_domain": False,
        },
        "implementation": {
            "module": "src/korea_business_lifecycle/authority_domain.py",
            "script": "scripts/authority_domain_reference.py",
            "reference_parser_dependency": "openpyxl==3.1.5",
            "current_snapshot_comparison_dependency": "pyarrow==21.0.0",
            "network_required_after_reference_download": False,
            "reference_workbook_git_tracked": False,
            "row_level_business_values_emitted": False,
        },
        "official_current_numeric_authorities": workbook["active_numeric_authorities"],
        "official_new_numeric_authorities": workbook["new_numeric_authorities"],
        "official_deleted_numeric_authorities": workbook["deleted_numeric_authorities"],
        "official_pre_reform_numeric_authorities": sorted(
            (
                {"code": code, "name": (current_map | deleted_map)[code]}
                for code in pre_reform
            ),
            key=lambda item: item["code"],
        ),
        "next_gate": (
            "use the approved date-effective current-state authority policy to choose one "
            "history cadence/request budget; API queryability of legacy/new partitions remains "
            "separate from date-effective membership"
        ),
    }
